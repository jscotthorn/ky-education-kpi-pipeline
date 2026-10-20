"""Kentucky Summative Assessment ETL module.

KYRC26 is two files. ``KYRC26_ACCT_Kentucky_Summative_Assessment.csv`` is
proficiency by school level. ``KYRC26_ASMT_Kentucky_Summative_Assessment.csv``
is proficiency by grade. The accountability headers match KYRC25. The
by-grade file changes co-op spelling only, and ``BaseETL`` already maps
both spellings.

Rates and the content index are stored as KDE publishes them. The
published proficient/distinguished percent is not the sum of the
proficient and distinguished percents, and a file with no combined
column does not invent one. A grade row is a grade metric.
It is not also stored as the elementary, middle, or high rate. Those
level rates come from the accountability file.

A ``*`` is value NA with ``suppressed`` Y. The star is read before the
cell is coerced. A blank cell is not stored. A published 0 is kept. A
row whose suppression flag is not Y and whose measure cells are all
blank was not offered and is skipped. One source row can mix a star
and a published content index, so the star suppresses only the starred
metrics.

``Assessment_Performance_by_Grade_2025.CSV`` is 404. The permanent
KYRC25 files are the 2024-25 source. A leftover temporary extract is
skipped. ``KYRC23`` republishes 2022-23 and is skipped.
"""
from __future__ import annotations

import logging
import resource
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from pydantic import BaseModel

etl_dir = Path(__file__).parent
sys.path.insert(0, str(etl_dir))

from base_etl import BaseETL
from constants import KPI_COLUMNS
from kyrc import normalize_kyrc_grade_token, release_year_from_filename

logger = logging.getLogger(__name__)


class Config(BaseModel):
    rename: Dict[str, str] = {}
    dtype: Dict[str, str] = {}
    derive: Dict[str, Any] = {}


# (normalized column, metric suffix, True when the cell is a 0-100 percent)
_MEASURES: Tuple[Tuple[str, str, bool], ...] = (
    ("novice", "novice_rate", True),
    ("apprentice", "apprentice_rate", True),
    ("proficient", "proficient_rate", True),
    ("distinguished", "distinguished_rate", True),
    ("proficient_distinguished", "proficient_distinguished_rate", True),
    ("content_index", "content_index_score", False),
)

_SUBJECTS = {
    "MA": "math",
    "Mathematics": "math",
    "RD": "reading",
    "Reading": "reading",
    "SC": "science",
    "Science": "science",
    "SS": "social_studies",
    "Social Studies": "social_studies",
    "WR": "writing",
    "Writing": "writing",
    "CW": "combined_writing",
    "Combined Writing": "combined_writing",
    "EM": "editing_and_mechanics",
    "Editing and Mechanics": "editing_and_mechanics",
    "On Demand Writing": "on_demand_writing",
    "On-Demand Writing": "on_demand_writing",
}

_LEVELS = {
    "ES": "elementary",
    "MS": "middle",
    "HS": "high",
    "Elementary School": "elementary",
    "Middle School": "middle",
    "High School": "high",
    "Elementary": "elementary",
    "Middle": "middle",
    "High": "high",
}

# Demographic labels on the 2017-18 and 2018-19 proficiency workbooks.
# The shared map drops ACD, sends CSG to Students with Disabilities, and
# sends LEN to English Learner. Those codes are the reverse of the
# 2017-18 coversheet. CSG is named on that coversheet only. The 2018-19
# coversheet does not list it, and the 2019 data still contains the code.
# It is still Consolidated Student Group. The labels below are applied
# before the shared mapper. The shared yaml is not edited.
_WORKBOOK_LABELS = {
    "ACD": "Students with Disabilities (IEP)",
    "ACO": "Students without IEP",
    "ACR": "Students with Disabilities/IEP Regular Assessment",
    "ACREG": "Students with Disabilities/IEP Regular Assessment",
    "ACY": "Students with Disabilities/IEP with Accommodations",
    "ACN": "Students with Disabilities/IEP without Accommodations",
    "CSG": "Consolidated Student Group",
    # CSN is not named on either coversheet. CSG + CSN equals all students,
    # so this is the complement of Consolidated Student Group, not students
    # without an IEP. The shared map sends CSN to Students without IEP.
    "CSN": "Not Consolidated Student Group",
    "LEN": "Non-English Learner",
    "ELN": "Non-English Learner or monitored",
    "MIN": "Non-Migrant",
}

# 2018-19 publishes these names on ``---COOP Total---`` rows. 2017-18
# publishes the same nine totals as ``---District Total---`` with a blank
# district name and school code 901-909. 909 has no abbreviation in 2019.
_COOP_TOTALS = {
    "901": ("Central Kentucky Educational Cooperative", "CKEC"),
    "902": ("Green River Regional Educational Cooperative", "GRREC"),
    "903": ("Kentucky Educational Development Corporation", "KEDC"),
    "904": ("Kentucky Valley Educational Corporation", "KVEC"),
    "905": ("Northern Kentucky Cooperative for Educational Services", "NKCES"),
    "906": ("Ohio Valley Educational Cooperative", "OVEC"),
    "907": ("Southeast/Southcentral Education Cooperative", "SESC"),
    "908": ("West Kentucky Educational Cooperative", "WKEC"),
    "909": ("Greater Louisville Educational Cooperative", None),
}

_BLANK_TOKENS = {"", "n/a", "na", "#n/a", "---", "--", "<10", "nan", "<na>", "none", "null"}

# Leftover temporary OAA names, and the KYRC23 republish of 2022-23.
_SKIP_FILENAMES = {
    "ASSESSMENT_PERFORMANCE_BY_GRADE_2024.CSV",
    "ASSESSMENT_PERFORMANCE_BY_GRADE_2025.CSV",
    "ASSESSMENT_PERFORMANCE_BY_GRADE_2026.CSV",
    "ACCOUNTABLE_ASSESSMENT_PERFORMANCE_2024.CSV",
    "ACCOUNTABLE_ASSESSMENT_PERFORMANCE_2025.CSV",
    "ACCOUNTABLE_ASSESSMENT_PERFORMANCE_2026.CSV",
    "KYRC23_ACCT_KENTUCKY_SUMMATIVE_ASSESSMENT.CSV",
    "KYRC23_ASMT_KENTUCKY_SUMMATIVE_ASSESSMENT.CSV",
}

_CHUNK_ROWS = 40000


def _raw_name(column: str) -> str:
    return f"__raw_{column}"


def _missing_scalar(value: Any) -> bool:
    """True for None and pandas NA. ``bool(pd.NA)`` raises, so this does not use it."""
    if value is None:
        return True
    try:
        missing = pd.isna(value)
    except (TypeError, ValueError):
        return False
    if isinstance(missing, bool):
        return missing
    return False


def _code_text(value: Any) -> str:
    """Identifier text. Blank, NA, and a trailing ``.0`` become a bare code."""
    if _missing_scalar(value):
        return ""
    text = str(value).strip()
    if text.lower() in {"", "nan", "<na>", "none", "null"}:
        return ""
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _classify_cell(value: Any, percent: bool) -> Tuple[str, Optional[float]]:
    """Return ``star``, ``number``, or ``blank`` for one measure cell.

    A star is recorded before the cell is coerced. A thousands comma is
    removed from the copied text. pandas 3 ``dtype=str`` is StringDtype,
    so this does not check ``dtype == object``.
    """
    if _missing_scalar(value):
        return "blank", None
    if isinstance(value, str):
        text = value.strip()
        if text in {"*", "**"}:
            return "star", None
        if text.lower() in _BLANK_TOKENS:
            return "blank", None
        cleaned = text.replace(",", "").replace("%", "")
        try:
            number = float(cleaned)
        except ValueError:
            return "blank", None
    else:
        try:
            number = float(value)
        except (TypeError, ValueError):
            return "blank", None
    if percent and (number < 0 or number > 100):
        return "blank", None
    if not percent and number < 0:
        return "blank", None
    return "number", number


def _classify_series(text: pd.Series, percent: bool) -> Tuple[pd.Series, pd.Series]:
    """Vector form of :func:`_classify_cell`."""
    raw = text.astype("string").str.strip()
    lower = raw.str.lower()
    star = raw.isin(["*", "**"]).fillna(False)
    blank_token = raw.isna() | raw.eq("") | lower.isin(list(_BLANK_TOKENS)).fillna(False)
    cleaned = raw.str.replace(",", "", regex=False).str.replace("%", "", regex=False)
    number = pd.to_numeric(cleaned, errors="coerce")
    number = number.mask(star | blank_token)
    if percent:
        invalid = number.notna() & ((number < 0) | (number > 100))
    else:
        invalid = number.notna() & (number < 0)
    if bool(invalid.any()):
        number = number.mask(invalid)
    kind = pd.Series("blank", index=text.index, dtype="object")
    kind = kind.mask(number.notna(), "number")
    kind = kind.mask(star, "star")
    return kind, number


class KentuckySummativeAssessmentETL(BaseETL):
    """ETL for Kentucky Summative Assessment proficiency percents."""

    # Opt in here, in the same change that accepts the KYRC26 headers.
    include_kyrc26 = True

    def accepts_data_file(self, file_path: Path) -> bool:
        """Skip a leftover temporary extract and the KYRC23 republish.

        ``Assessment_Performance_by_Grade_2025.CSV`` and
        ``Accountable_Assessment_Performance_2025.CSV`` are 404. The
        2026 temporary names still return 200 and use 3-digit codes.
        A leftover copy would write 2024, 2025, or 2026 a second time.
        ``KYRC23`` is school year 20222023, the year the registered
        2023 files already store. The match is case-insensitive.
        """
        if file_path.name.upper() in _SKIP_FILENAMES:
            logger.info(
                "Skipping %s; a permanent file already stores that year",
                file_path.name,
            )
            return False
        return super().accepts_data_file(file_path)

    @property
    def module_column_mappings(self) -> Dict[str, str]:
        return {
            "Level": "level",
            "LEVEL": "level",
            "Subject": "subject",
            "SUBJECT": "subject",
            "Novice": "novice",
            "NOVICE": "novice",
            "Apprentice": "apprentice",
            "APPRENTICE": "apprentice",
            "Proficient": "proficient",
            "PROFICIENT": "proficient",
            "Distinguished": "distinguished",
            "DISTINGUISHED": "distinguished",
            "distinguished": "distinguished",
            "Proficient / Distinguished": "proficient_distinguished",
            "Proficient/Distinguished": "proficient_distinguished",
            "PROFICIENT/DISTINGUISHED": "proficient_distinguished",
            "PROFICIENT_DISTINGUISHED": "proficient_distinguished",
            "Content Index": "content_index",
            "CONTENT INDEX": "content_index",
            "SCH_YEAR": "school_year",
            "CNTYNO": "county_number",
            "CNTYNAME": "county_name",
            "DIST_NUMBER": "district_number",
            "DIST_NAME": "district_name",
            "SCH_NUMBER": "school_number",
            "SCH_NAME": "school_name",
            "SCH_CD": "school_code",
            "STATE_SCH_ID": "state_school_id",
            "NCESID": "nces_id",
            "COOP": "co_op",
            "COOP_CODE": "co_op_code",
            "GRADE": "grade",
            "TESTED": "tested",
        }

    def standardize_missing_values(self, df: pd.DataFrame) -> pd.DataFrame:
        """Keep a copy of each measure cell, including stars.

        ``BaseETL.standardize_missing_values`` only rewrites
        ``dtype == object``. pandas 3 ``dtype=str`` is StringDtype, so a
        star would otherwise stay in the column on one pandas and become
        NA on the other. The copy is the source of the suppression flag.
        The base method is not changed.
        """
        for column, _suffix, _percent in _MEASURES:
            if column in df.columns:
                df[_raw_name(column)] = df[column].astype("string")
        return super().standardize_missing_values(df)

    def _normalize_subject(self, value: Any) -> str:
        if _missing_scalar(value) or str(value).strip() == "":
            return "unknown_subject"
        text = str(value).strip()
        return _SUBJECTS.get(text, text.lower().replace(" ", "_"))

    def _normalize_grade(self, grade: Any) -> Optional[str]:
        """Grade period, or None when the row is not a grade row.

        ``3rd`` and ``10th`` become ``grade_3`` and ``grade_10`` through
        :func:`normalize_kyrc_grade_token`. ``03`` and ``Grade 3`` already
        did. A grade row is not also given a level period.
        """
        if _missing_scalar(grade):
            return None
        text = str(grade).strip()
        if text.lower() in _BLANK_TOKENS or text.lower() in {"all_grades", "all grades"}:
            return None
        mapped = normalize_kyrc_grade_token(text).strip()
        if mapped.lower() in _BLANK_TOKENS:
            return None
        lowered = mapped.lower()
        if lowered.startswith("grade_"):
            rest = lowered.split("_", 1)[1]
            if rest.isdigit():
                return f"grade_{int(rest)}"
            return lowered
        if mapped.isdigit():
            return f"grade_{int(mapped)}"
        if lowered.startswith("grade "):
            rest = mapped[6:].strip()
            if rest.isdigit():
                return f"grade_{int(rest)}"
            return "grade_" + rest.lower().replace(" ", "_")
        return lowered.replace(" ", "_")

    def _normalize_level(self, level: Any) -> Optional[str]:
        if _missing_scalar(level) or str(level).strip() == "":
            return None
        text = str(level).strip()
        if text.lower() in _BLANK_TOKENS or text.lower() == "all":
            return None
        mapped = _LEVELS.get(text)
        if mapped:
            return mapped
        lowered = text.lower().replace(" ", "_")
        if lowered == "all":
            return None
        return lowered

    def _periods_for_row(self, row: pd.Series) -> List[str]:
        """Grade period when the row has a grade, otherwise the level period.

        The 2018-19 grade workbook carries both columns. The level cell
        is the school's level, and the percent is the grade percent.
        Storing that percent as elementary would duplicate the level file.
        """
        grade_period = self._normalize_grade(row.get("grade"))
        if grade_period:
            return [grade_period]
        level_period = self._normalize_level(row.get("level"))
        if level_period:
            return [level_period]
        return []

    def _measure_text(self, row: pd.Series, column: str) -> Any:
        raw = _raw_name(column)
        if raw in row.index:
            return row.get(raw)
        if column in row.index:
            return row.get(column)
        return pd.NA

    def _row_measures(self, row: pd.Series) -> List[Tuple[str, str, bool]]:
        present = []
        for column, suffix, percent in _MEASURES:
            if _raw_name(column) in row.index or column in row.index:
                present.append((column, suffix, percent))
        return present

    def _row_is_blank(self, row: pd.Series) -> bool:
        """True when every present measure cell is empty, not a star or a 0."""
        measures = self._row_measures(row)
        if not measures:
            return True
        for column, _suffix, percent in measures:
            kind, _number = _classify_cell(self._measure_text(row, column), percent)
            if kind != "blank":
                return False
        return True

    def _apply_workbook_labels(self, df: pd.DataFrame, source_file: str) -> pd.DataFrame:
        """Read proficiency-workbook codes from the coversheet.

        Applied only to ``ASSESSMENT_PROFICIENCY_*.xlsx``. Other years
        already use the labels the shared map expects.
        """
        if not Path(source_file).name.upper().startswith("ASSESSMENT_PROFICIENCY_"):
            return df
        if "demographic" not in df.columns:
            return df
        labels = df["demographic"].astype("string").str.strip().map(_WORKBOOK_LABELS)
        present = labels.notna()
        if not bool(present.any()):
            return df
        out = df.copy()
        out.loc[present, "demographic"] = labels.loc[present]
        return out

    def _label_workbook_coop_totals(self, frame: pd.DataFrame, source_file: str) -> pd.DataFrame:
        """Label 2017-18 co-op totals the way 2019 publishes them.

        Districts 901-909 in the 2017-18 proficiency workbooks are
        ``---District Total---`` with a blank district name. Stored that
        way they look like nine unnamed districts. 2019 already publishes
        the long co-op name, ``---COOP Total---``, and the abbreviation
        (blank for 909). A 2019 row has a district name, so this does not
        rewrite it. Only ``ASSESSMENT_PROFICIENCY_*`` files are labelled.
        """
        if not Path(source_file).name.upper().startswith("ASSESSMENT_PROFICIENCY_"):
            return frame
        if "district_number" not in frame.columns or "school_name" not in frame.columns:
            return frame
        codes = frame["district_number"].map(_code_text)
        if "district_name" in frame.columns:
            blank_district = frame["district_name"].map(_code_text).eq("")
        else:
            blank_district = pd.Series(True, index=frame.index)
        school = frame["school_name"].map(
            lambda value: "" if _missing_scalar(value) else str(value).strip()
        )
        mask = codes.isin(_COOP_TOTALS) & blank_district & school.eq("---District Total---")
        if not bool(mask.any()):
            return frame
        out = frame.copy()
        for column in ("district_name", "co_op", "co_op_code"):
            if column not in out.columns:
                out[column] = pd.Series(pd.NA, index=out.index, dtype="object")

        def long_name(code: str) -> Any:
            if code not in _COOP_TOTALS:
                return pd.NA
            return _COOP_TOTALS[code][0]

        def abbreviation(code: str) -> Any:
            if code not in _COOP_TOTALS:
                return pd.NA
            published = _COOP_TOTALS[code][1]
            return published if published else pd.NA

        selected = codes.loc[mask]
        out.loc[mask, "district_name"] = selected.map(long_name)
        out.loc[mask, "school_name"] = "---COOP Total---"
        out.loc[mask, "co_op_code"] = selected
        out.loc[mask, "co_op"] = selected.map(abbreviation)
        return out

    def should_skip_row(self, row: pd.Series) -> bool:
        """Skip a blank, not-suppressed row. A star is not blank."""
        if super().should_skip_row(row):
            return True
        suppressed = row.get("suppressed", "N")
        if _missing_scalar(suppressed):
            flag = ""
        else:
            flag = str(suppressed).strip()
        if flag == "Y":
            return False
        return self._row_is_blank(row)

    def extract_metrics(self, row: pd.Series) -> Dict[str, Any]:
        """Published percents for one row.

        A star is stored as NA so the caller can mark that metric
        suppressed. A blank cell is omitted. There is no proficient plus
        distinguished fallback when the combined column is absent.
        """
        subject = self._normalize_subject(row.get("subject"))
        periods = self._periods_for_row(row)
        metrics: Dict[str, Any] = {}
        if not periods:
            return metrics
        for column, suffix, percent in self._row_measures(row):
            kind, number = _classify_cell(self._measure_text(row, column), percent)
            if kind == "blank":
                continue
            value: Any = pd.NA if kind == "star" else number
            for period in periods:
                metrics[f"kentucky_summative_assessment_{subject}_{suffix}_{period}"] = value
        return metrics

    def get_suppressed_metric_defaults(self, row: pd.Series) -> Dict[str, Any]:
        """NA for each metric a fully suppressed row would have stored."""
        subject = self._normalize_subject(row.get("subject"))
        periods = self._periods_for_row(row)
        metrics: Dict[str, Any] = {}
        for _column, suffix, _percent in self._row_measures(row):
            for period in periods:
                metrics[f"kentucky_summative_assessment_{subject}_{suffix}_{period}"] = pd.NA
        return metrics

    def _map_group(self, value: Any, year: str, source_file: str, cache: Dict[Tuple[str, str], Optional[str]]) -> Optional[str]:
        key = ("", year) if _missing_scalar(value) else (str(value).strip(), year)
        if key not in cache:
            cache[key] = self.demographic_mapper.map_demographic(
                None if key[0] == "" else key[0],
                year,
                source_file,
            )
        return cache[key]

    def convert_to_kpi_format(self, df: pd.DataFrame, source_file: str) -> pd.DataFrame:
        """One KPI row per published cell, or per star.

        A grade row writes grade metrics only. A level row writes level
        metrics only. Blank cells are left out. A star on a row that also
        publishes a content index suppresses the star only.
        """
        if df.empty:
            return pd.DataFrame()
        frame = self._apply_workbook_labels(df, source_file)
        frame = self._label_workbook_coop_totals(frame, source_file)
        if "demographic" not in frame.columns:
            return pd.DataFrame()

        demographic = frame["demographic"]
        keep_demo = demographic.notna()
        if "demographic" in frame.columns:
            demo_text = demographic.astype("string").str.strip()
            keep_demo = keep_demo & ~demo_text.isin(self.DISTRICT_AGGREGATE_PATTERNS).fillna(False)
        frame = frame.loc[keep_demo]
        if frame.empty:
            return pd.DataFrame()

        years = self._years(frame, source_file)
        cache: Dict[Tuple[str, str], Optional[str]] = {}
        groups: List[Optional[str]] = []
        for value, year in zip(frame["demographic"].tolist(), years.tolist()):
            groups.append(self._map_group(value, str(year), source_file, cache))
        group_series = pd.Series(groups, index=frame.index, dtype="object")
        frame = frame.loc[group_series.notna()].copy()
        group_series = group_series.loc[group_series.notna()]
        years = years.loc[frame.index]
        if frame.empty:
            return pd.DataFrame()

        subjects = frame["subject"].map(self._normalize_subject) if "subject" in frame.columns else pd.Series("unknown_subject", index=frame.index)
        grade_period = self._grade_periods(frame)
        level_period = self._level_periods(frame)
        has_grade = grade_period.notna()
        period = grade_period.where(has_grade, level_period)
        usable = period.notna()
        frame = frame.loc[usable]
        if frame.empty:
            return pd.DataFrame()
        group_series = group_series.loc[frame.index]
        years = years.loc[frame.index]
        subjects = subjects.loc[frame.index]
        period = period.loc[frame.index]

        kinds: Dict[str, pd.Series] = {}
        numbers: Dict[str, pd.Series] = {}
        present_measures = []
        for column, suffix, percent in _MEASURES:
            raw = self._series_text(frame, column)
            if raw is None:
                continue
            kind, number = _classify_series(raw, percent)
            kinds[column] = kind
            numbers[column] = number
            present_measures.append((column, suffix, percent))
        if not present_measures:
            return pd.DataFrame()

        all_blank = pd.Series(True, index=frame.index)
        for column, _suffix, _percent in present_measures:
            all_blank = all_blank & kinds[column].eq("blank")
        suppressed = self._suppressed_flag(frame)
        skip = all_blank & suppressed.ne("Y")
        frame = frame.loc[~skip]
        if frame.empty:
            return pd.DataFrame()
        group_series = group_series.loc[frame.index]
        years = years.loc[frame.index]
        subjects = subjects.loc[frame.index]
        period = period.loc[frame.index]
        all_blank = all_blank.loc[frame.index]
        suppressed = suppressed.loc[frame.index]
        for column in list(kinds):
            kinds[column] = kinds[column].loc[frame.index]
            numbers[column] = numbers[column].loc[frame.index]

        school_id = self._school_ids(frame)
        school_name = self._school_names(frame["school_name"]) if "school_name" in frame.columns else pd.Series("Unknown School", index=frame.index)
        district = frame["district_name"] if "district_name" in frame.columns else pd.Series("Unknown District", index=frame.index)
        now = pd.Timestamp.now().isoformat()

        pieces: List[pd.DataFrame] = []
        base_columns = {
            "year": years.astype("string"),
            "district": district.astype("string"),
            "school_name": school_name.astype("string"),
            "student_group": group_series.astype("string"),
            "school_id": school_id.astype("string"),
            "source_file": source_file,
            "last_updated": now,
        }
        for attribute in (
            "county_number",
            "county_name",
            "district_number",
            "school_code",
            "state_school_id",
            "nces_id",
            "co_op",
            "co_op_code",
            "school_type",
        ):
            if attribute in frame.columns:
                base_columns[attribute] = frame[attribute]

        for column, suffix, _percent in present_measures:
            kind = kinds[column]
            number = numbers[column]
            metric = "kentucky_summative_assessment_" + subjects.astype("string") + "_" + suffix + "_" + period.astype("string")
            star_or_number = kind.isin(["star", "number"])
            if bool(star_or_number.any()):
                piece = pd.DataFrame(base_columns).loc[star_or_number].copy()
                piece["metric"] = metric.loc[star_or_number].to_numpy()
                piece["value"] = number.loc[star_or_number].to_numpy()
                piece["suppressed"] = pd.Series(
                    ["Y" if flag == "star" else "N" for flag in kind.loc[star_or_number].tolist()],
                    index=piece.index,
                )
                pieces.append(piece)
            # A fully blank Y row has no star to carry the suppression.
            fill = all_blank & suppressed.eq("Y")
            if bool(fill.any()):
                piece = pd.DataFrame(base_columns).loc[fill].copy()
                piece["metric"] = metric.loc[fill].to_numpy()
                piece["value"] = pd.NA
                piece["suppressed"] = "Y"
                pieces.append(piece)

        if not pieces:
            return pd.DataFrame()
        for piece in pieces:
            piece["value"] = pd.to_numeric(piece["value"], errors="coerce")
        kpi = pd.concat(pieces, ignore_index=True)
        ordered = [column for column in KPI_COLUMNS if column in kpi.columns]
        return kpi[ordered]

    def _series_text(self, frame: pd.DataFrame, column: str) -> Optional[pd.Series]:
        raw = _raw_name(column)
        if raw in frame.columns:
            return frame[raw]
        if column in frame.columns:
            return frame[column]
        return None

    def _years(self, frame: pd.DataFrame, source_file: str) -> pd.Series:
        fallback = release_year_from_filename(source_file) or "2024"
        if "school_year" not in frame.columns:
            return pd.Series(fallback, index=frame.index, dtype="string")
        text = frame["school_year"].astype("string").str.strip()
        last4 = text.str[-4:]
        use_last4 = text.str.len().eq(8) | text.str.len().eq(4)
        years = last4.where(use_last4, fallback)
        years = years.mask(text.isna() | text.eq("") | text.str.lower().isin(["<na>", "nan", "none"]), fallback)
        return years.astype("string")

    def _grade_periods(self, frame: pd.DataFrame) -> pd.Series:
        if "grade" not in frame.columns:
            return pd.Series(pd.NA, index=frame.index, dtype="object")
        return frame["grade"].map(self._normalize_grade)

    def _level_periods(self, frame: pd.DataFrame) -> pd.Series:
        if "level" not in frame.columns:
            return pd.Series(pd.NA, index=frame.index, dtype="object")
        return frame["level"].map(self._normalize_level)

    def _suppressed_flag(self, frame: pd.DataFrame) -> pd.Series:
        if "suppressed" not in frame.columns:
            return pd.Series("N", index=frame.index, dtype="string")
        flag = frame["suppressed"].astype("string").str.strip()
        flag = flag.mask(flag.isna() | flag.str.lower().isin(["<na>", "nan", "none", ""]), "")
        return flag

    def _school_ids(self, frame: pd.DataFrame) -> pd.Series:
        """School code, except a co-op total is keyed by 901-909.

        ``ASSESSMENT_PROFICIENCY_*_18-19`` publishes ``---COOP Total---``
        with school code ``000`` for each co-op. The published co-op code,
        or the district number when that cell is blank, is the key.
        ``_label_workbook_coop_totals`` renames the 2017-18 rows, which
        already use school code 901-909, before this runs. A school whose
        name is not exactly ``---COOP Total---`` keeps its school code.
        """
        if "school_code" not in frame.columns:
            raise ValueError("No valid school ID found in row")
        text = frame["school_code"].map(_code_text)
        if bool(text.eq("").any()):
            logger.error("CRITICAL: No valid school ID found in row")
            raise ValueError("No valid school ID found in row")
        if "school_name" not in frame.columns:
            return text
        school = frame["school_name"].map(
            lambda value: "" if _missing_scalar(value) else str(value).strip()
        )
        coop = school.eq("---COOP Total---")
        if not bool(coop.any()):
            return text
        if "co_op_code" in frame.columns:
            coop_code = frame["co_op_code"].map(_code_text)
        else:
            coop_code = pd.Series("", index=frame.index)
        if "district_number" in frame.columns:
            district_code = frame["district_number"].map(_code_text)
        else:
            district_code = pd.Series("", index=frame.index)
        chosen = coop_code.where(coop_code.isin(_COOP_TOTALS), district_code)
        chosen = chosen.where(chosen.isin(_COOP_TOTALS), text)
        return text.where(~coop, chosen)

    def _school_names(self, series: pd.Series) -> pd.Series:
        text = series.astype("string").str.strip()
        blank = series.isna() | text.isna() | text.eq("") | text.str.lower().isin(["<na>", "nan", "none"])
        return text.mask(blank | text.eq("All Schools"), "---District Total---")

    def _prepare_frame(self, df: pd.DataFrame, source_file: str, conf: Config) -> pd.DataFrame:
        df = self.normalize_column_names(df)
        if "grade" in df.columns:
            df["grade"] = df["grade"].map(
                lambda value: normalize_kyrc_grade_token(value) if not _missing_scalar(value) else value
            )
        df = self.standardize_missing_values(df)
        df = self.normalize_grade_field(df)
        df = self.add_derived_fields(df, conf.derive, source_file)
        if conf.rename:
            df = df.rename(columns=conf.rename)
        if conf.dtype:
            for col, dtype in conf.dtype.items():
                if col not in df.columns:
                    continue
                try:
                    if (
                        dtype.lower().startswith("float")
                        or dtype.lower().startswith("int")
                        or dtype in ["Int64", "Int32", "Float64"]
                    ):
                        df[col] = pd.to_numeric(df[col], errors="coerce")
                    df[col] = df[col].astype(dtype)
                except Exception as exc:
                    logger.warning("Failed to convert column %s to %s: %s", col, dtype, exc)
        return df

    def _write_kpi(
        self,
        kpi: pd.DataFrame,
        writer_state: Dict[str, Any],
        demographic_tracker: Dict[str, set],
    ) -> int:
        if kpi.empty:
            return 0
        columns = list(KPI_COLUMNS)
        out = kpi.reindex(columns=columns)
        out.to_csv(
            writer_state["handle"],
            index=False,
            header=not writer_state["started"],
            na_rep="",
        )
        writer_state["started"] = True
        pairs = out.loc[
            out["year"].notna() & out["student_group"].notna(),
            ["year", "student_group"],
        ].drop_duplicates()
        for year, group in pairs.itertuples(index=False):
            year_text = "" if _missing_scalar(year) else str(year)
            group_text = "" if _missing_scalar(group) else str(group)
            if year_text and group_text:
                demographic_tracker.setdefault(year_text, set()).add(group_text)
        if hasattr(self, "_year_rows"):
            years = out["year"].astype("string")
            for year, count in years.value_counts(dropna=False).items():
                key = "" if _missing_scalar(year) else str(year)
                self._year_rows[key] = self._year_rows.get(key, 0) + int(count)
            starred = out.loc[out["suppressed"].astype("string").eq("Y"), "year"].astype("string")
            for year, count in starred.value_counts(dropna=False).items():
                key = "" if _missing_scalar(year) else str(year)
                self._year_suppressed[key] = self._year_suppressed.get(key, 0) + int(count)
        return len(out)

    def process(self, raw_dir: Path, proc_dir: Path, cfg: dict) -> None:
        """Stream each file in chunks so the proficiency tables stay under 4 GB."""
        source_dir = raw_dir / self.source_name
        if not source_dir.exists():
            logger.info("No raw data directory for %s; skipping.", self.source_name)
            return
        data_files = self._find_data_files(source_dir)
        if not data_files:
            logger.info("No data files found in %s; skipping.", source_dir)
            return

        output_path = proc_dir / f"{self.source_name}.csv"
        conf = Config(**cfg)
        total = 0
        demographic_tracker: Dict[str, set] = {}
        self._year_rows: Dict[str, int] = {}
        self._year_suppressed: Dict[str, int] = {}
        started = time.perf_counter()
        with output_path.open("w", newline="", encoding="utf-8") as handle:
            writer_state: Dict[str, Any] = {"handle": handle, "writer": None, "started": False, "columns": []}
            for data_file, sheet in data_files:
                logger.info("Processing %s", data_file.name)
                try:
                    if data_file.stat().st_size == 0:
                        logger.warning("Empty file (0 bytes): %s", data_file.name)
                        continue
                    file_rows = self._process_one_file(
                        data_file, sheet, conf, writer_state, demographic_tracker
                    )
                    total += file_rows
                    logger.info("Finished %s: %s KPI rows", data_file.name, f"{file_rows:,}")
                except Exception as exc:
                    logger.error("Error processing %s: %s", data_file.name, exc)
                    continue
        if total == 0:
            logger.warning("No valid KPI data files processed")
            return
        validation = self._validate_demographics_from_tracker(demographic_tracker)
        self._save_demographic_report(proc_dir, validation)
        elapsed = time.perf_counter() - started
        peak_kb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        logger.info(
            "Completed %s: %s KPI rows in %.1fs, peak rss %.1f MB. "
            "Rows by year: %s. Suppressed Y by year: %s",
            self.source_name,
            f"{total:,}",
            elapsed,
            peak_kb / 1024,
            dict(sorted(self._year_rows.items())),
            dict(sorted(self._year_suppressed.items())),
        )
        print(f"Wrote {output_path}")
        print(f"Demographic report: {proc_dir / f'{self.source_name}_demographic_report.md'}")

    def _process_one_file(
        self,
        data_file: Path,
        sheet: Optional[str],
        conf: Config,
        writer_state: Dict[str, Any],
        demographic_tracker: Dict[str, set],
    ) -> int:
        written = 0
        if data_file.suffix.lower() == ".csv":
            # na_filter is off so a star and a literal N/A stay text on
            # pandas 2 and pandas 3. The classifier reads that text.
            reader = pd.read_csv(
                data_file,
                encoding="utf-8-sig",
                dtype=str,
                chunksize=_CHUNK_ROWS,
                keep_default_na=False,
                na_filter=False,
            )
            for chunk in reader:
                prepared = self._prepare_frame(chunk, data_file.name, conf)
                kpi = self.convert_to_kpi_format(prepared, data_file.name)
                written += self._write_kpi(kpi, writer_state, demographic_tracker)
                del prepared, kpi, chunk
            return written

        for chunk in self._iter_xlsx_chunks(data_file, sheet):
            prepared = self._prepare_frame(chunk, data_file.name, conf)
            kpi = self.convert_to_kpi_format(prepared, data_file.name)
            written += self._write_kpi(kpi, writer_state, demographic_tracker)
            del prepared, kpi, chunk
        return written

    def _iter_xlsx_chunks(self, data_file: Path, sheet: Optional[str]):
        """Yield string chunks from one workbook sheet.

        The grade workbooks are large. Reading the whole sheet with
        pandas holds it at once. openpyxl walks the sheet and this
        method keeps one chunk.
        """
        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.error("openpyxl is required to read %s", data_file.name)
            return
        wb = load_workbook(data_file, read_only=True, data_only=True)
        try:
            chosen = sheet or "DATA"
            if chosen not in wb.sheetnames:
                logger.error(
                    "Sheet %s not in %s (sheets: %s)",
                    chosen,
                    data_file.name,
                    wb.sheetnames,
                )
                return
            iterator = wb[chosen].iter_rows(values_only=True)
            header_row = next(iterator, None)
            if not header_row:
                return
            header = ["" if cell is None else str(cell).strip() for cell in header_row]
            width = len(header)
            buffer: List[List[str]] = []
            for raw in iterator:
                if raw is None:
                    continue
                values: List[str] = []
                for index in range(width):
                    cell = raw[index] if index < len(raw) else None
                    values.append("" if cell is None else str(cell).strip())
                if not any(values):
                    continue
                buffer.append(values)
                if len(buffer) >= _CHUNK_ROWS:
                    yield pd.DataFrame(buffer, columns=header)
                    buffer = []
            if buffer:
                yield pd.DataFrame(buffer, columns=header)
        finally:
            wb.close()


def transform(raw_dir: Path, proc_dir: Path, cfg: dict) -> None:
    """Entry point for pipeline."""
    etl = KentuckySummativeAssessmentETL("kentucky_summative_assessment")
    etl.process(raw_dir, proc_dir, cfg)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    root = Path(__file__).parent.parent
    raw_dir = root / "data" / "raw"
    proc_dir = root / "data" / "processed"
    proc_dir.mkdir(exist_ok=True)
    transform(raw_dir, proc_dir, {"derive": {}})
